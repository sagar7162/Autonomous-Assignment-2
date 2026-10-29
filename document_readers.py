"""Extract local documents with stable, human-checkable source locations."""
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SUPPORTED = {'.pdf', '.html', '.htm', '.docx', '.xlsx'}


def office_pages(path):
    path = Path(path)
    if path.suffix.lower() == '.docx':
        document = Document(path)
        paragraphs = '\n'.join(f'Paragraph {i}: {p.text}' for i, p in enumerate(document.paragraphs, 1) if p.text.strip())
        pages = [('document paragraphs', paragraphs, False)]
        for i, table in enumerate(document.tables, 1):
            header = ' | '.join(cell.text for cell in table.rows[0].cells) if table.rows else ''
            for j, row in enumerate(table.rows, 1):
                text = f'{path.stem}; Table {i}; Header: {header}; Row {j}: ' + ' | '.join(cell.text for cell in row.cells)
                pages.append((f'table {i}, row {j}', text, False))
        return pages
    if path.suffix.lower() == '.xlsx':
        workbook = load_workbook(path, data_only=True, read_only=False)
        pages = []
        try:
            for sheet in workbook:
                # Repeat initial headings so isolated rows retain column meanings.
                rows = []
                for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    values = ' | '.join(f'{get_column_letter(i)}={str(v).strip()}' for i, v in enumerate(row, 1) if v is not None and str(v).strip())
                    if values:
                        rows.append((number, values))
                headings = '; '.join(f'row {i}: {text}' for i, text in rows[:3])
                for number, values in rows:
                    text = f'{path.stem}; Sheet: {sheet.title}; Headings: {headings}\nRow {number}: {values}'
                    pages.append((f'sheet {sheet.title}, row {number}', text, False))
        finally:
            workbook.close()
        return pages
    raise ValueError(f'Unsupported Office document: {path.name}')


def html_pages(path):
    soup = BeautifulSoup(Path(path).read_bytes(), 'html.parser')
    for node in soup(['script', 'style', 'nav', 'footer', 'header']):
        node.decompose()
    return [('web', soup.get_text(' ', strip=True), False)]
