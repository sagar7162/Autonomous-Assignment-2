import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import chromadb
from docx import Document
from openpyxl import Workbook

import helpdesk
from document_readers import office_pages
from register_documents import register


class IngestionTests(unittest.TestCase):
    def test_registration_is_idempotent_and_skips_identical_files(self):
        with tempfile.TemporaryDirectory(dir=helpdesk.ROOT) as temp:
            root = Path(temp)
            (root / 'documents').mkdir()
            (root / 'documents/manifest.json').write_text('[]', encoding='utf-8')
            (root / 'documents/one.pdf').write_bytes(b'example source')
            (root / 'documents/copy.pdf').write_bytes(b'example source')
            first = register(root)
            self.assertEqual(len(first['added_source_ids']), 1)
            self.assertEqual(len(first['duplicates']), 1)
            self.assertEqual(register(root)['added_source_ids'], [])
            manifest = json.loads((root / 'documents/manifest.json').read_text())
            self.assertEqual(len(manifest), 1)
            self.assertTrue(manifest[0]['url'].startswith('file:'))

    def test_spreadsheet_preserves_sheet_row_and_sparse_columns(self):
        with tempfile.TemporaryDirectory(dir=helpdesk.ROOT) as temp:
            path = Path(temp) / 'semester.xlsx'
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = 'Classroom Allocation'
            sheet.append(['Course Code', 'Subject', None, 'Slot', 'Instructor'])
            sheet.append(['CSL7650', 'Autonomous Systems', None, 'E', 'Gaurav Harit'])
            workbook.save(path)
            workbook.close()
            pages = office_pages(path)
            location, text, ocr = pages[1]
            self.assertEqual(location, 'sheet Classroom Allocation, row 2')
            for expected in ('CSL7650', 'Autonomous Systems', 'D=E', 'Gaurav Harit', 'Slot'):
                self.assertIn(expected, text)
            self.assertFalse(ocr)

    def test_word_includes_policy_text_and_table(self):
        with tempfile.TemporaryDirectory(dir=helpdesk.ROOT) as temp:
            path = Path(temp) / 'guidelines.docx'
            document = Document()
            document.add_paragraph('Students must show an institute identity card.')
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = 'Rule'
            table.cell(0, 1).text = 'Requirement'
            table.cell(1, 0).text = 'Arrival'
            table.cell(1, 1).text = 'Ten minutes before the examination'
            document.save(path)
            pages = office_pages(path)
            self.assertIn('identity card', pages[0][1])
            self.assertTrue(any('Ten minutes' in text and location == 'table 1, row 2' for location, text, _ in pages))

    def test_failed_rebuild_preserves_live_collection(self):
        with tempfile.TemporaryDirectory(dir=helpdesk.ROOT) as temp:
            root = Path(temp)
            (root / 'documents').mkdir()
            (root / 'documents/source.html').write_text('<p>' + 'Academic regulation. ' * 20 + '</p>')
            (root / 'documents/manifest.json').write_text(json.dumps([{'id': 'test', 'title': 'Test', 'file': 'documents/source.html', 'url': 'https://example.edu', 'status': 'downloaded'}]))
            client = chromadb.EphemeralClient()
            live = client.create_collection('academic')
            live.add(ids=['original'], documents=['original evidence'], embeddings=[[1.0, 0.0]])
            try:
                with patch.object(helpdesk, 'ROOT', root), patch.object(helpdesk.chromadb, 'PersistentClient', return_value=client), patch.object(helpdesk, 'embeddings', side_effect=RuntimeError('embedding service failed')):
                    with self.assertRaisesRegex(RuntimeError, 'embedding service failed'):
                        helpdesk.ingest()
                self.assertEqual(client.get_collection('academic').get()['ids'], ['original'])
            finally:
                for collection in client.list_collections():
                    client.delete_collection(collection.name)

    def test_successful_rebuild_replaces_old_chunks(self):
        with tempfile.TemporaryDirectory(dir=helpdesk.ROOT) as temp:
            root = Path(temp)
            (root / 'documents').mkdir()
            (root / 'documents/source.html').write_text('<p>' + 'New academic regulation. ' * 20 + '</p>')
            (root / 'documents/manifest.json').write_text(json.dumps([{'id': 'test', 'title': 'Test', 'file': 'documents/source.html', 'url': 'https://example.edu', 'status': 'downloaded'}]))
            client = chromadb.EphemeralClient()
            client.create_collection('academic').add(ids=['original'], embeddings=[[1.0, 0.0]])
            try:
                with patch.object(helpdesk, 'ROOT', root), patch.object(helpdesk.chromadb, 'PersistentClient', return_value=client), patch.object(helpdesk, 'embeddings', side_effect=lambda batch: [[0.0, 1.0] for _ in batch]):
                    helpdesk.ingest()
                result = client.get_collection('academic').get(include=['documents', 'metadatas'])
                self.assertNotIn('original', result['ids'])
                self.assertEqual(result['metadatas'][0]['source_id'], 'test')
                self.assertIn('New academic regulation', result['documents'][0])
            finally:
                for collection in client.list_collections():
                    client.delete_collection(collection.name)


if __name__ == '__main__':
    unittest.main()
